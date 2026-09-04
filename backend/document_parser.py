"""
Multi-format document parser with server-side MIME verification, input sanitization, and security checks.
Supports: PDF, DOCX, TXT, MD, HTML, CSV, XLSX, PNG, JPG, JPEG, BMP, TIFF, and Web URLs
"""

import os
import re
import io
import zipfile
import logging
import socket
import ipaddress
from urllib.parse import urlparse
from typing import Tuple, Dict, Any, Optional

logger = logging.getLogger(__name__)


def is_safe_url(url: str) -> bool:
    """Validate that a URL does not resolve to loopback, private, or cloud metadata addresses."""
    try:
        parsed = urlparse(url)
        if parsed.scheme not in ("http", "https"):
            return False
        hostname = parsed.hostname
        if not hostname:
            return False
        if hostname.lower() in ("localhost", "127.0.0.1", "::1", "0.0.0.0"):
            return False
        addr_info = socket.getaddrinfo(hostname, None)
        for entry in addr_info:
            ip_str = entry[4][0]
            ip = ipaddress.ip_address(ip_str)
            if (
                ip.is_loopback
                or ip.is_private
                or ip.is_link_local
                or ip.is_reserved
                or ip.is_multicast
                or str(ip) in ("169.254.169.254", "0.0.0.0", "255.255.255.255")
            ):
                return False
        return True
    except Exception:
        return False


class DocumentParser:
    SUPPORTED_TYPES = {
        ".pdf", ".docx", ".txt", ".md", ".html", ".htm",
        ".csv", ".xlsx", ".png", ".jpg", ".jpeg", ".bmp", ".tiff"
    }

    # ──────────────────────────────────────────────
    # Server-Side Input Sanitization & Verification
    # ──────────────────────────────────────────────

    @classmethod
    def sanitize_filename(cls, raw_filename: str) -> str:
        """Sanitize raw filenames to prevent path traversal and shell injection."""
        base = os.path.basename(raw_filename or "document.txt")
        # Remove null bytes and path separators
        base = base.replace("\x00", "").replace("/", "").replace("\\", "")
        # Replace non-safe characters with underscore
        clean = re.sub(r"[^\w.\-]", "_", base)
        # Collapse multiple dots or underscores
        clean = re.sub(r"\.{2,}", ".", clean)
        clean = re.sub(r"_{2,}", "_", clean)
        # Prevent hidden files
        if clean.startswith("."):
            clean = f"upload_{clean}"
        if not clean or clean == "_":
            clean = "upload_document.txt"
        # Truncate to reasonable length (100 chars) while preserving extension
        if len(clean) > 100:
            name, ext = os.path.splitext(clean)
            clean = f"{name[:90]}{ext}"
        return clean

    @classmethod
    def validate_upload(
        cls,
        file_bytes: bytes,
        raw_filename: str,
        max_size_bytes: int = 52_428_800
    ) -> Tuple[str, str]:
        """
        Server-side validation:
        1. Checks file size
        2. Sanitizes filename
        3. Verifies file extension is supported
        4. Verifies actual binary magic bytes / MIME signatures
        5. Protects against zip bombs, executable headers, and malicious payloads
        Returns: (clean_filename, extension)
        """
        # 1. Size verification
        if not file_bytes or len(file_bytes) == 0:
            raise ValueError("Uploaded file is empty (0 bytes).")
        if len(file_bytes) > max_size_bytes:
            max_mb = max_size_bytes / (1024 * 1024)
            actual_mb = len(file_bytes) / (1024 * 1024)
            raise ValueError(f"File size ({actual_mb:.1f} MB) exceeds maximum allowed limit of {max_mb:.0f} MB.")

        # 2. Filename sanitization
        clean_filename = cls.sanitize_filename(raw_filename)
        ext = os.path.splitext(clean_filename)[1].lower()

        # 3. Extension check
        if ext not in cls.SUPPORTED_TYPES:
            raise ValueError(
                f"Unsupported file type '{ext}'. Supported formats: {', '.join(sorted(cls.SUPPORTED_TYPES))}"
            )

        # 4. Binary signature (magic bytes) verification
        if ext == ".pdf":
            # PDF must contain %PDF- near beginning
            if not file_bytes.startswith(b"%PDF-") and b"%PDF-" not in file_bytes[:1024]:
                raise ValueError("File failed MIME verification: Invalid PDF header signature.")

        elif ext == ".docx":
            if not file_bytes.startswith(b"PK\x03\x04"):
                raise ValueError("File failed MIME verification: Invalid DOCX signature.")
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                    namelist = zf.namelist()
                    if not any("word/document.xml" in n or "[Content_Types].xml" in n for n in namelist):
                        raise ValueError("File failed verification: Corrupted or invalid Word DOCX package.")
                    # Decompression bomb prevention
                    total_uncompressed = sum(info.file_size for info in zf.infolist())
                    if total_uncompressed > max_size_bytes * 10:
                        raise ValueError("File rejected: Suspected decompression bomb in DOCX.")
            except zipfile.BadZipFile:
                raise ValueError("File failed verification: Corrupted ZIP/DOCX archive.")

        elif ext == ".xlsx":
            if not file_bytes.startswith(b"PK\x03\x04"):
                raise ValueError("File failed MIME verification: Invalid XLSX signature.")
            try:
                with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                    namelist = zf.namelist()
                    if not any("xl/workbook.xml" in n or "[Content_Types].xml" in n for n in namelist):
                        raise ValueError("File failed verification: Corrupted or invalid Excel XLSX package.")
                    total_uncompressed = sum(info.file_size for info in zf.infolist())
                    if total_uncompressed > max_size_bytes * 10:
                        raise ValueError("File rejected: Suspected decompression bomb in XLSX.")
            except zipfile.BadZipFile:
                raise ValueError("File failed verification: Corrupted ZIP/XLSX archive.")

        elif ext in (".png",):
            if not file_bytes.startswith(b"\x89PNG\r\n\x1a\n"):
                raise ValueError("File failed MIME verification: Invalid PNG signature.")

        elif ext in (".jpg", ".jpeg"):
            if not file_bytes.startswith(b"\xff\xd8\xff"):
                raise ValueError("File failed MIME verification: Invalid JPEG signature.")

        elif ext == ".bmp":
            if not file_bytes.startswith(b"BM"):
                raise ValueError("File failed MIME verification: Invalid BMP signature.")

        elif ext == ".tiff":
            if not (file_bytes.startswith(b"II*\x00") or file_bytes.startswith(b"MM\x00*")):
                raise ValueError("File failed MIME verification: Invalid TIFF signature.")

        elif ext in (".txt", ".md", ".csv", ".html", ".htm"):
            # Check for executable headers (Windows PE / Linux ELF / Mach-O / Java Class)
            if file_bytes.startswith(b"MZ") or file_bytes.startswith(b"\x7fELF") or file_bytes.startswith(b"\xca\xfe\xba\xbe"):
                raise ValueError("Executable binary files cannot be uploaded as plain text.")
            # Check for excessive null bytes (indicative of binary/compiled payload)
            null_count = file_bytes[:4096].count(b"\x00")
            if null_count > 10:
                raise ValueError("File appears to be a binary executable, not valid plain text.")

        return clean_filename, ext

    # ──────────────────────────────────────────────
    # Parser Entrypoints
    # ──────────────────────────────────────────────

    def parse(self, file_path: str, filename: str) -> Tuple[str, Dict[str, Any]]:
        """
        Parse a document and return (full_text, metadata).
        metadata includes: title, source, page_count, word_count, file_type
        """
        ext = os.path.splitext(filename)[1].lower()

        if ext not in self.SUPPORTED_TYPES:
            raise ValueError(
                f"Unsupported file type '{ext}'. "
                f"Supported: {', '.join(sorted(self.SUPPORTED_TYPES))}"
            )

        if ext == ".pdf":
            return self._parse_pdf(file_path, filename)
        elif ext == ".docx":
            return self._parse_docx(file_path, filename)
        elif ext in (".txt", ".md"):
            return self._parse_text(file_path, filename)
        elif ext in (".html", ".htm"):
            return self._parse_html(file_path, filename)
        elif ext == ".csv":
            return self._parse_csv(file_path, filename)
        elif ext == ".xlsx":
            return self._parse_xlsx(file_path, filename)
        elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".tiff"):
            return self._parse_image(file_path, filename)

        raise ValueError(f"Unhandled file type: {ext}")

    def parse_url(self, url: str) -> Tuple[str, Dict[str, Any]]:
        """
        Fetch and parse a public Web URL safely.
        """
        import urllib.request
        from bs4 import BeautifulSoup

        # Enforce HTTP/HTTPS only
        if not url.startswith("http://") and not url.startswith("https://"):
            raise ValueError("URL must start with http:// or https://")

        if not is_safe_url(url):
            raise ValueError("Access to local, private, or loopback network addresses is prohibited.")

        req = urllib.request.Request(
            url,
            headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) DocMind/2.0"}
        )

        try:
            with urllib.request.urlopen(req, timeout=15) as resp:
                # Limit URL response size to 10 MB
                html = resp.read(10_485_760).decode("utf-8", errors="ignore")
        except Exception as exc:
            raise ValueError(f"Failed to fetch URL '{url}': {exc}")

        soup = BeautifulSoup(html, "html.parser")
        # Sanitize HTML tags completely
        for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript", "iframe", "object", "embed"]):
            tag.decompose()

        title = soup.title.string.strip() if soup.title and soup.title.string else url

        parts = []
        for tag in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th"]):
            text = tag.get_text(strip=True)
            if not text:
                continue
            if tag.name in ("h1", "h2", "h3", "h4", "h5", "h6"):
                level = int(tag.name[1])
                parts.append(f"{'#' * level} {text}")
            else:
                parts.append(text)

        full_text = "\n\n".join(parts)
        full_text = self._clean_text(full_text)
        if not full_text:
            full_text = f"(Web page '{url}' contained no readable main text.)"

        metadata = {
            "title": title,
            "source": url,
            "file_type": "url",
            "page_count": 1,
            "word_count": len(full_text.split()),
        }
        return full_text, metadata

    # ──────────────────────────────────────────────
    # Specific Parsers
    # ──────────────────────────────────────────────

    def _parse_pdf(self, path: str, filename: str) -> Tuple[str, Dict]:
        import fitz  # pymupdf

        doc = fitz.open(path)
        try:
            pages_text = []
            for i, page in enumerate(doc, 1):
                text = page.get_text("text")
                if text and text.strip():
                    pages_text.append(f"--- [Page {i}] ---\n{text.strip()}")

            full_text = "\n\n".join(pages_text)
            full_text = self._clean_text(full_text)

            # If no extractable text, attempt OCR fallback on PDF pages
            if not full_text or len(full_text.split()) < 10:
                ocr_text = self._ocr_pdf_pages(doc, filename)
                if ocr_text and not ocr_text.startswith("(Document '"):
                    full_text = ocr_text
                elif not full_text:
                    full_text = ocr_text

            metadata = {
                "title": doc.metadata.get("title") or filename,
                "source": filename,
                "file_type": "pdf",
                "page_count": len(doc),
                "word_count": len(full_text.split()),
            }
            return full_text, metadata
        finally:
            doc.close()

    def _ocr_pdf_pages(self, doc, filename: str) -> str:
        """OCR fallback for scanned PDF files."""
        try:
            import pytesseract
            from PIL import Image

            ocr_pages = []
            for i, page in enumerate(doc, 1):
                pix = page.get_pixmap(dpi=150)
                img = Image.open(io.BytesIO(pix.tobytes()))
                text = pytesseract.image_to_string(img)
                if text.strip():
                    ocr_pages.append(f"--- [Page {i}] ---\n{text.strip()}")

            result = "\n\n".join(ocr_pages)
            if result.strip():
                return self._clean_text(result)
        except Exception as exc:
            logger.warning("OCR processing failed for '%s': %s", filename, exc)

        return f"(Document '{filename}' appears to contain scanned images or no extractable text.)"

    def _parse_image(self, path: str, filename: str) -> Tuple[str, Dict]:
        """Parse image files via OCR."""
        text = ""
        try:
            import pytesseract
            from PIL import Image

            img = Image.open(path)
            text = pytesseract.image_to_string(img)
        except Exception as exc:
            logger.warning("Image OCR failed for '%s': %s", filename, exc)

        cleaned = self._clean_text(text)
        if not cleaned:
            cleaned = f"(Image '{filename}' contains no readable OCR text or Tesseract OCR is not configured.)"

        metadata = {
            "title": filename,
            "source": filename,
            "file_type": os.path.splitext(filename)[1].lstrip(".").lower(),
            "page_count": 1,
            "word_count": len(cleaned.split()),
        }
        return cleaned, metadata

    def _parse_docx(self, path: str, filename: str) -> Tuple[str, Dict]:
        from docx import Document

        doc = Document(path)
        sections = []

        for para in doc.paragraphs:
            style_name = getattr(para.style, "name", "") or ""
            if style_name.startswith("Heading"):
                level = style_name.split()[-1]
                prefix = "#" * int(level) if level.isdigit() else "##"
                sections.append(f"{prefix} {para.text.strip()}")
            elif para.text.strip():
                sections.append(para.text.strip())

        # Include tables
        for table in doc.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            if rows:
                sections.append("\n".join(rows))

        full_text = "\n\n".join(sections)
        full_text = self._clean_text(full_text)
        if not full_text:
            full_text = f"(Document '{filename}' contains no extractable paragraph text.)"

        metadata = {
            "title": filename,
            "source": filename,
            "file_type": "docx",
            "page_count": None,
            "word_count": len(full_text.split()),
        }
        return full_text, metadata

    def _parse_text(self, path: str, filename: str) -> Tuple[str, Dict]:
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
        except Exception:
            with open(path, "r", encoding="latin-1", errors="replace") as f:
                text = f.read()

        text = self._clean_text(text)
        if not text:
            text = f"(File '{filename}' is empty.)"

        metadata = {
            "title": filename,
            "source": filename,
            "file_type": "txt" if filename.endswith(".txt") else "md",
            "page_count": None,
            "word_count": len(text.split()),
        }
        return text, metadata

    def _parse_html(self, path: str, filename: str) -> Tuple[str, Dict]:
        from bs4 import BeautifulSoup

        try:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                html = f.read()
        except Exception:
            with open(path, "r", encoding="latin-1", errors="replace") as f:
                html = f.read()

        soup = BeautifulSoup(html, "html.parser")

        # Strip scripts, styles, iframes, objects
        for tag in soup(["script", "style", "nav", "footer", "header", "iframe", "object", "embed", "noscript"]):
            tag.decompose()

        title = soup.title.string if soup.title else filename

        parts = []
        for tag in soup.find_all(
            ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th"]
        ):
            text = tag.get_text(strip=True)
            if not text:
                continue
            if tag.name in ("h1", "h2", "h3", "h4", "h5", "h6"):
                level = int(tag.name[1])
                parts.append(f"{'#' * level} {text}")
            else:
                parts.append(text)

        full_text = "\n\n".join(parts)
        full_text = self._clean_text(full_text)
        if not full_text:
            full_text = f"(Web page '{filename}' contains no readable textual content.)"

        metadata = {
            "title": title,
            "source": filename,
            "file_type": "html",
            "page_count": None,
            "word_count": len(full_text.split()),
        }
        return full_text, metadata

    def _parse_csv(self, path: str, filename: str) -> Tuple[str, Dict]:
        import pandas as pd

        try:
            df = pd.read_csv(path, nrows=5000)
        except UnicodeDecodeError:
            df = pd.read_csv(path, encoding="latin-1", nrows=5000)

        # Sanitize CSV formula injection for preview
        def sanitize_cell(v: Any) -> str:
            s = str(v)
            if s.startswith(("=", "+", "-", "@", "\t", "\r")):
                return "'" + s
            return s

        display_df = df.head(500)
        rows = [" | ".join(sanitize_cell(v) for v in display_df.columns)]
        rows.append("-" * len(rows[0]))
        for _, row in display_df.iterrows():
            rows.append(" | ".join(sanitize_cell(v) for v in row.values))

        data_text = "\n".join(rows)
        try:
            summary = df.describe(include="all").to_string()
        except Exception:
            summary = "Summary unavailable."

        full_text = f"# CSV Data: {filename}\n\n## Data (Preview)\n{data_text}\n\n## Summary\n{summary}"
        full_text = self._clean_text(full_text)

        metadata = {
            "title": filename,
            "source": filename,
            "file_type": "csv",
            "page_count": None,
            "word_count": len(full_text.split()),
            "row_count": len(df),
            "col_count": len(df.columns),
        }
        return full_text, metadata

    def _parse_xlsx(self, path: str, filename: str) -> Tuple[str, Dict]:
        import pandas as pd

        def sanitize_cell(v: Any) -> str:
            s = str(v) if v is not None else ""
            if s.startswith(("=", "+", "-", "@", "\t", "\r")):
                return "'" + s
            return s

        xls = pd.ExcelFile(path)
        try:
            parts = []
            for sheet_name in xls.sheet_names[:10]:
                df = pd.read_excel(xls, sheet_name=sheet_name, nrows=500)
                rows = [" | ".join(sanitize_cell(v) for v in df.columns)]
                rows.append("-" * len(rows[0]))
                for _, row in df.iterrows():
                    rows.append(" | ".join(sanitize_cell(v) for v in row.values))
                parts.append(f"## Sheet: {sheet_name}\n\n" + "\n".join(rows))

            full_text = f"# Excel File: {filename}\n\n" + "\n\n".join(parts)
            full_text = self._clean_text(full_text)

            metadata = {
                "title": filename,
                "source": filename,
                "file_type": "xlsx",
                "page_count": len(xls.sheet_names),
                "word_count": len(full_text.split()),
            }
            return full_text, metadata
        finally:
            xls.close()

    # ──────────────────────────────────────────────
    # Helpers
    # ──────────────────────────────────────────────

    @staticmethod
    def _clean_text(text: str) -> str:
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        # Normalize docx bullet symbols (e.g. \uf0b7, •, ▪, ‣) to standard '- '
        text = re.sub(r"^[ \t]*[\uf0b7•▪‣⁃][ \t]*", "- ", text, flags=re.MULTILINE)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = text.replace("\x00", "")
        return text.strip()
